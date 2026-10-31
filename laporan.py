"""Laporan siap cetak (PDF dan Word) dari data berita SBSN.

Susunan: 1. Ringkasan (kartu KPI + grafik, tata letak seperti dashboard) · 2. Insight utama ·
3. Pembahasan dan insight per topik · Lampiran daftar berita (opsional).

Dipakai oleh app.py. Bisa juga dari terminal:
    python laporan.py hasil/berita_sbsn_2026-10.xlsx            # PDF + Word
"""

import io
import re
import sys
from collections import Counter
from datetime import datetime
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402

try:
    import scraper_sbsn_chrome as _sc
    _POLA_POS, _POLA_NEG = _sc._POLA_POS, _sc._POLA_NEG
except Exception:  # laporan tetap bisa dibuat tanpa modul scraper
    _POLA_POS = _POLA_NEG = None

UNGU, EMAS, MAGENTA, PLUM = "#6C007D", "#FFC000", "#D0208A", "#7C3F84"
TONE = ["Positif", "Netral", "Negatif"]
WARNA_TONE = {"Positif": "#2FA888", "Netral": "#3B82C4", "Negatif": "#E0533D"}
TINT_TONE = {"Positif": "#D9F2E3", "Netral": "#DCEAF7", "Negatif": "#F9D9D9"}
PALET_TOPIK = [UNGU, EMAS, MAGENTA, "#002060", PLUM, "#2FA888", "#B4C7E7"]
NAMA_BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
              "Agustus", "September", "Oktober", "November", "Desember"]
UNIT = "Direktorat Pembiayaan Syariah · DJPPR Kementerian Keuangan"

STOPWORDS = set("""
yang dan di ke dari untuk dengan ini itu pada akan dalam jadi bisa ada tak tidak rp
hingga sebagai atau juga oleh karena lebih sudah masih cara apa begini berikut simak
mulai hari kini baru bagi para pun serta agar saat usai jelang telah tersebut tahun
bulan pekan the of to in and sbsn ri vs per capai buka tahap via kata soal triliun miliar
""".split())


def tgl_id(t, pendek=False):
    if t is None or pd.isna(t):
        return "-"
    b = NAMA_BULAN[t.month - 1]
    return f"{t.day} {b[:3] if pendek else b} {t.year}"


def _pct(x):
    return f"{x:.0%}".replace(".", ",")


# ===========================================================================
# ANALISIS
# ===========================================================================

def _siapkan(df):
    df = df.copy()
    for k in ["Tanggal", "Situs", "Judul", "Link", "Topik", "Tone", "Ritel"]:
        if k not in df.columns:
            df[k] = ""
    df["Tanggal"] = pd.to_datetime(df["Tanggal"], errors="coerce")
    for k in ["Situs", "Judul", "Link", "Topik", "Tone", "Ritel"]:
        df[k] = df[k].fillna("").astype(str)
    df["Topik"] = df["Topik"].replace("", "(tanpa topik)")
    # Urutan baku: tanggal terlama -> terbaru (berlaku juga untuk lampiran daftar berita)
    return df.sort_values("Tanggal", ascending=True, na_position="last", kind="stable").reset_index(drop=True)


def kata_teratas(judul, n=10, kecuali=()):
    kecuali = {k.lower() for k in kecuali}
    c = Counter()
    for j in judul:
        for w in re.findall(r"[A-Za-z][A-Za-z0-9]+", j):
            lw = w.lower()
            if len(lw) > 2 and lw not in STOPWORDS and lw not in kecuali:
                c[w.upper() if re.fullmatch(r"(SR|ST|SW)\d{3}(T\d+)?", w, re.I) else lw] += 1
    return c.most_common(n)


def _pemicu(judul, pola, n=4):
    if pola is None:
        return []
    c = Counter(m.group(1).lower() for j in judul for m in pola.finditer(j))
    return [k for k, _ in c.most_common(n)]


def _indeks(tone_counts, n):
    """Indeks sentimen -100..100 = (positif - negatif) / total."""
    if not n:
        return 0
    return round((tone_counts.get("Positif", 0) - tone_counts.get("Negatif", 0)) / n * 100)


def _label_indeks(x):
    if x >= 30:
        return "sangat positif"
    if x >= 10:
        return "cenderung positif"
    if x > -10:
        return "berimbang"
    if x > -30:
        return "cenderung negatif"
    return "negatif"


def _perubahan(kini, lalu):
    if lalu == 0:
        return None
    return (kini - lalu) / lalu


def _seri_hitung(df):
    s = df["Ritel"].str.split(", ").explode()
    s = s[s.fillna("") != ""]
    return s.value_counts()


def analisis_topik(nama, d, total, d_lalu=None, nama_lalu=None):
    n = len(d)
    tc = d["Tone"].value_counts().to_dict()
    idx = _indeks(tc, n)
    a = dict(nama=nama, n=n, share=n / total if total else 0, tone={t: int(tc.get(t, 0)) for t in TONE},
             indeks=idx, label=_label_indeks(idx))

    kata_topik = re.findall(r"[A-Za-z]+", nama)
    a["kata"] = [w for w, _ in kata_teratas(d["Judul"], 6, kecuali=kata_topik)]
    a["pemicu_pos"] = _pemicu(d.loc[d["Tone"] == "Positif", "Judul"], _POLA_POS)
    a["pemicu_neg"] = _pemicu(d.loc[d["Tone"] == "Negatif", "Judul"], _POLA_NEG)
    media = d["Situs"].replace("", pd.NA).dropna().value_counts()
    a["media"] = list(media.head(3).items())
    a["konsentrasi"] = (media.iloc[0] / n) if n and len(media) else 0
    harian = d.dropna(subset=["Tanggal"]).groupby(d["Tanggal"].dt.normalize()).size()
    a["puncak"] = (harian.idxmax(), int(harian.max())) if len(harian) and harian.max() >= 3 else None
    a["seri"] = list(_seri_hitung(d).head(5).items())
    a["contoh_pos"] = d[d["Tone"] == "Positif"].head(2)[["Judul", "Situs", "Link"]].values.tolist()
    a["contoh_neg"] = d[d["Tone"] == "Negatif"].head(2)[["Judul", "Situs", "Link"]].values.tolist()

    ada_lalu = d_lalu is not None
    n_lalu = len(d_lalu) if ada_lalu else 0
    a["n_lalu"] = n_lalu if ada_lalu else None
    a["ubah"] = _perubahan(n, n_lalu) if ada_lalu else None
    a["indeks_lalu"] = (_indeks(d_lalu["Tone"].value_counts().to_dict(), n_lalu)
                        if ada_lalu and n_lalu else None)

    # ---- Pembahasan (narasi) ----
    kal = [f"Topik {nama} muncul dalam {n} berita ({_pct(a['share'])} dari seluruh pemberitaan)"]
    if ada_lalu:
        if n_lalu == 0:
            kal[0] += f", padahal pada {nama_lalu} tidak tercatat pemberitaan"
        elif a["ubah"] is not None and abs(a["ubah"]) >= 0.05:
            kal[0] += (f", {'naik' if a['ubah'] > 0 else 'turun'} {_pct(abs(a['ubah']))} "
                       f"dibanding {nama_lalu} ({n_lalu} berita)")
        else:
            kal[0] += f", relatif stabil dibanding {nama_lalu} ({n_lalu} berita)"
    kal[0] += "."
    kal.append(f"Nada pemberitaan {a['label']} (indeks sentimen {idx:+d}) dengan komposisi "
               f"{a['tone']['Positif']} positif, {a['tone']['Netral']} netral, dan "
               f"{a['tone']['Negatif']} negatif.")
    if a["kata"]:
        kal.append("Isu yang paling sering diangkat di judul berita antara lain: "
                   + ", ".join(a["kata"][:5]) + ".")
    if a["media"]:
        kal.append("Media yang paling aktif memberitakan: "
                   + ", ".join(f"{m} ({k})" for m, k in a["media"]) + ".")
    if a["puncak"]:
        kal.append(f"Puncak pemberitaan terjadi pada {tgl_id(a['puncak'][0])} "
                   f"dengan {a['puncak'][1]} berita.")
    a["pembahasan"] = " ".join(kal)

    # ---- Insight ----
    ins = []
    if a["pemicu_pos"]:
        ins.append("Nada positif terutama dipicu oleh isu " + ", ".join(f"“{k}”" for k in a["pemicu_pos"]) + ".")
    if a["pemicu_neg"]:
        ins.append("Nada negatif terutama terkait isu " + ", ".join(f"“{k}”" for k in a["pemicu_neg"]) + ".")
    if a["indeks_lalu"] is not None:
        sel = idx - a["indeks_lalu"]
        if abs(sel) >= 10:
            ins.append(f"Sentimen {'membaik' if sel > 0 else 'memburuk'} dibanding {nama_lalu} "
                       f"(indeks {a['indeks_lalu']:+d} menjadi {idx:+d}).")
        else:
            ins.append(f"Sentimen relatif sama dengan {nama_lalu} (indeks {a['indeks_lalu']:+d} menjadi {idx:+d}).")
    if a["konsentrasi"] >= 0.4 and n >= 5:
        ins.append(f"Pemberitaan terkonsentrasi di {a['media'][0][0]} ({_pct(a['konsentrasi'])}); "
                   "jangkauan ke media lain masih terbatas.")
    elif len(a["media"]) >= 3 and n >= 5:
        ins.append("Pemberitaan tersebar di banyak media, menandakan jangkauan isu yang luas.")
    if a["seri"]:
        ins.append("Seri SBSN Ritel yang diberitakan: "
                   + ", ".join(f"{s} ({k} berita)" for s, k in a["seri"]) + ".")
    if not ins:
        ins.append("Belum ada pola menonjol; jumlah berita masih sedikit.")
    a["insight"] = ins

    # ---- Rekomendasi ----
    rek = []
    neg_share = a["tone"]["Negatif"] / n if n else 0
    pos_share = a["tone"]["Positif"] / n if n else 0
    if neg_share >= 0.3:
        isu = (" terkait isu " + ", ".join(f"“{k}”" for k in a["pemicu_neg"][:2])) if a["pemicu_neg"] else ""
        rek.append(f"Siapkan klarifikasi atau penguatan narasi{isu} melalui siaran pers, "
                   "media briefing, atau konten penjelasan di kanal resmi.")
    if pos_share >= 0.5:
        isu = (" (" + ", ".join(f"“{k}”" for k in a["pemicu_pos"][:2]) + ")") if a["pemicu_pos"] else ""
        rek.append(f"Pertahankan momentum positif{isu}; amplifikasi capaian lewat media sosial "
                   "dan mitra distribusi.")
    if a["tone"]["Netral"] / (n or 1) >= 0.5:
        rek.append("Pemberitaan cenderung informatif; tambahkan sudut pandang manfaat dan dampak "
                   "agar pesan lebih kuat.")
    if a["ubah"] is not None and a["ubah"] <= -0.3:
        rek.append("Eksposur menurun; pertimbangkan rilis berita, wawancara, atau konten edukasi "
                   "untuk menjaga visibilitas.")
    if a["konsentrasi"] >= 0.4 and n >= 5:
        rek.append("Perluas jejaring media agar pemberitaan tidak bergantung pada satu media.")
    if not rek:
        rek.append("Lanjutkan pemantauan rutin; belum ada sinyal yang memerlukan tindakan khusus.")
    a["rekomendasi"] = rek
    return a


def analisis(df, periode, df_lalu=None, nama_lalu=None):
    df = _siapkan(df)
    ada_lalu = df_lalu is not None and len(df_lalu) > 0
    dl = _siapkan(df_lalu) if ada_lalu else None
    n = len(df)
    tc = df["Tone"].value_counts().to_dict()
    A = dict(df=df, periode=periode, nama_lalu=nama_lalu, ada_lalu=ada_lalu, n=n,
             n_lalu=len(dl) if ada_lalu else None,
             tone={t: int(tc.get(t, 0)) for t in TONE},
             tone_lalu=({t: int(dl["Tone"].value_counts().get(t, 0)) for t in TONE} if ada_lalu else None),
             indeks=_indeks(tc, n))

    urutan = df["Topik"].value_counts().index.tolist()
    A["topik"] = [analisis_topik(t, df[df["Topik"] == t], n,
                                 dl[dl["Topik"] == t] if ada_lalu else None, nama_lalu)
                  for t in urutan]

    # Ringkasan naratif
    kal = [f"Selama periode {periode} tercatat {n} berita terkait SBSN"]
    if ada_lalu:
        d = n - A["n_lalu"]
        kal[0] += (f", {'naik' if d > 0 else 'turun'} {abs(d)} berita dibanding {nama_lalu}"
                   if d else f", sama dengan {nama_lalu}")
    kal[0] += "."
    if n:
        kal.append(f"Secara keseluruhan nada pemberitaan {_label_indeks(A['indeks'])} "
                   f"(indeks sentimen {A['indeks']:+d}): {A['tone']['Positif']} positif, "
                   f"{A['tone']['Netral']} netral, {A['tone']['Negatif']} negatif.")
        kal.append("Sebaran topik: " + ", ".join(f"{t['nama']} {t['n']}" for t in A["topik"]) + ".")
        ms = df["Situs"].replace("", pd.NA).dropna().value_counts().head(3)
        if len(ms):
            kal.append("Media yang paling banyak memberitakan: "
                       + ", ".join(f"{s} ({k})" for s, k in ms.items()) + ".")
    A["ringkasan"] = " ".join(kal)

    # Insight utama lintas topik
    iu = []
    if A["topik"]:
        dom = A["topik"][0]
        iu.append(f"{dom['nama']} menjadi topik paling banyak diberitakan ({dom['n']} berita, "
                  f"{_pct(dom['share'])}).")
        terbaik = max(A["topik"], key=lambda t: (t["indeks"], t["n"]))
        terburuk = min(A["topik"], key=lambda t: (t["indeks"], -t["n"]))
        if terbaik["indeks"] >= 10:
            iu.append(f"Sentimen paling positif ada pada {terbaik['nama']} (indeks {terbaik['indeks']:+d})"
                      + (f", didorong isu {', '.join(terbaik['pemicu_pos'][:2])}." if terbaik["pemicu_pos"] else "."))
        if terburuk["indeks"] <= -10 and terburuk is not terbaik:
            iu.append(f"{terburuk['nama']} perlu perhatian karena sentimennya {terburuk['label']} "
                      f"(indeks {terburuk['indeks']:+d})"
                      + (f", terutama isu {', '.join(terburuk['pemicu_neg'][:2])}." if terburuk["pemicu_neg"] else "."))
        naik = [t for t in A["topik"] if t["ubah"] is not None and t["ubah"] >= 0.3]
        turun = [t for t in A["topik"] if t["ubah"] is not None and t["ubah"] <= -0.3]
        if naik:
            iu.append("Volume pemberitaan meningkat signifikan pada "
                      + ", ".join(f"{t['nama']} (+{_pct(t['ubah'])})" for t in naik) + ".")
        if turun:
            iu.append("Volume pemberitaan menurun pada "
                      + ", ".join(f"{t['nama']} (-{_pct(abs(t['ubah']))})" for t in turun) + ".")
        seri = _seri_hitung(df)
        if len(seri):
            iu.append(f"Seri ritel yang paling banyak diberitakan: {seri.index[0]} ({int(seri.iloc[0])} berita).")
    A["insight_utama"] = iu or ["Belum ada data yang cukup untuk ditarik insight."]

    seri = df.assign(Seri=df["Ritel"].str.split(", ")).explode("Seri").reset_index(drop=True)
    seri = seri[seri["Seri"].fillna("") != ""]
    A["rekap_seri"] = (pd.crosstab(seri["Seri"], seri["Tone"]).reindex(columns=TONE, fill_value=0)
                       .assign(Total=lambda x: x.sum(axis=1)).sort_values("Total", ascending=False)
                       if len(seri) else None)
    A["media"] = df["Situs"].replace("", pd.NA).dropna().value_counts().head(10)
    A["negatif"] = df[df["Tone"] == "Negatif"]
    return A


# ===========================================================================
# GRAFIK (matplotlib -> PNG bytes)
# ===========================================================================

def _png_tetap(fig, kiri=0.1, kanan=0.98, atas=0.95, bawah=0.2):
    """Simpan gambar dengan ukuran PERSIS seperti figsize (tanpa dipotong), supaya panel yang
    berdampingan bisa dibuat sama tinggi."""
    fig.subplots_adjust(left=kiri, right=kanan, top=atas, bottom=bawah)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def rentang_tampil(df, rentang=None):
    """Rentang tanggal sumbu tren: rentang yang diberikan, atau seluruh bulan data
    (tanggal 1 sampai akhir bulan)."""
    if rentang and rentang[0] is not None:
        return pd.Timestamp(rentang[0]), pd.Timestamp(rentang[1])
    t = pd.to_datetime(df["Tanggal"], errors="coerce").dropna()
    if t.empty:
        return None, None
    awal = t.min().replace(day=1)
    akhir = (t.max().replace(day=1) + pd.offsets.MonthEnd(0))
    return awal.normalize(), akhir.normalize()


def _png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _donut(label, nilai, warna, judul):
    fig, ax = plt.subplots(figsize=(4.2, 3.0))
    total = sum(nilai)
    if total == 0:
        ax.text(0.5, 0.5, "Tidak ada data", ha="center", va="center")
        ax.axis("off")
        return _png(fig)
    w, _ = ax.pie(nilai, colors=warna, startangle=90, counterclock=False,
                  wedgeprops=dict(width=0.36, edgecolor="white", linewidth=1.5))
    ax.text(0, 0.06, f"{total}", ha="center", va="center", fontsize=20, fontweight="bold", color="#2A2A2A")
    ax.text(0, -0.2, "berita", ha="center", va="center", fontsize=9, color="#6B6B6B")
    ax.legend(w, [f"{l}  {v} ({v / total:.0%})" for l, v in zip(label, nilai)],
              loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, fontsize=8.5)
    ax.set_title(judul, fontsize=10.5, fontweight="bold", color=UNGU, loc="left")
    ax.set_aspect("equal")
    return _png(fig)


def png_donut_tone(df):
    h = df["Tone"].value_counts().reindex(TONE, fill_value=0)
    return _donut(TONE, list(h.values), [WARNA_TONE[t] for t in TONE], "Komposisi tone")


def png_donut_topik(df):
    h = df["Topik"].value_counts()
    return _donut(list(h.index), list(h.values),
                  [PALET_TOPIK[i % len(PALET_TOPIK)] for i in range(len(h))], "Berita per topik")


def png_tren(df):
    d = df.dropna(subset=["Tanggal"])
    if d.empty:
        return None
    tabel = (d.groupby([d["Tanggal"].dt.normalize(), "Tone"]).size().unstack(fill_value=0)
             .reindex(columns=TONE, fill_value=0))
    tabel = tabel.reindex(pd.date_range(tabel.index.min(), tabel.index.max()), fill_value=0)
    fig, ax = plt.subplots(figsize=(8.6, 2.6))
    bawah = None
    for t in TONE:
        ax.bar(tabel.index, tabel[t], bottom=bawah, color=WARNA_TONE[t], label=t, width=0.8)
        bawah = tabel[t] if bawah is None else bawah + tabel[t]
    ax.set_title("Tren harian jumlah berita", fontsize=10.5, fontweight="bold", color=UNGU, loc="left")
    ax.xaxis.set_major_formatter(matplotlib.dates.DateFormatter("%d/%m"))
    ax.yaxis.get_major_locator().set_params(integer=True)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#E6E1EA")
    ax.set_axisbelow(True)
    ax.tick_params(labelsize=8)
    ax.legend(frameon=False, fontsize=8, ncol=3, loc="lower right", bbox_to_anchor=(1, 1.0))
    return _png(fig)


def png_tone_topik(t):
    """Bar horizontal kecil komposisi tone untuk satu topik."""
    fig, ax = plt.subplots(figsize=(6.5, 0.55))
    kiri = 0
    for tone in TONE:
        v = t["tone"][tone]
        if v:
            ax.barh([0], [v], left=kiri, color=D_TONE[tone], height=0.8)
            ax.text(kiri + v / 2, 0, f"{tone} {v}", ha="center", va="center", fontsize=7.5,
                    color="white", fontweight="bold")
            kiri += v
    ax.set_xlim(0, max(kiri, 1))
    ax.axis("off")
    return _png(fig)


# ---------------------------------------------------------------------------
# Grafik bergaya dashboard (warna & susunan seperti dashboard_sbsn.html)
# ---------------------------------------------------------------------------
D_POS, D_NET, D_NEG = "#23977A", "#3B82C4", "#D24A35"   # Netral = biru
D_TONE = {"Positif": D_POS, "Netral": D_NET, "Negatif": D_NEG}
D_UNGU, D_EMAS, D_FG, D_MUTED, D_LINE = "#6C007D", "#FFC000", "#26202B", "#6E6574", "#E3DCE8"
BLN3 = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]


def _gaya_ax(ax, grid_axis="y"):
    for s in ["top", "right", "left", "bottom"]:
        ax.spines[s].set_visible(False)
    ax.grid(axis=grid_axis, color=D_LINE, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(colors=D_MUTED, labelsize=7.5, length=0)


def _legenda(ax, ncol=3):
    from matplotlib.lines import Line2D
    h = [Line2D([0], [0], marker="o", color="w", markerfacecolor=D_TONE[t], markersize=6, label=t)
         for t in TONE]
    ax.legend(handles=h, loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=ncol, frameon=False,
              fontsize=7.5, labelcolor=D_MUTED, handletextpad=0.3, columnspacing=1.2)


def _cm(w_cm, h_cm):
    return plt.subplots(figsize=(w_cm / 2.54, h_cm / 2.54))


def _kiri_cm(cm_, w_cm):
    return min(0.45, cm_ / w_cm)


def png_dash_tren(df, w_cm, h_cm, rentang=None):
    """Tren harian per tone, sumbu-x berisi SETIAP tanggal dari awal sampai akhir rentang."""
    fig, ax = _cm(w_cm, h_cm)
    awal, akhir = rentang_tampil(df, rentang)
    d = df.dropna(subset=["Tanggal"])
    if awal is None:
        ax.text(0.5, 0.5, "Tidak ada data bertanggal", ha="center", va="center", color=D_MUTED)
        ax.axis("off")
        return _png_tetap(fig)
    hari = pd.date_range(awal, akhir, freq="D")
    tabel = (d.groupby([d["Tanggal"].dt.normalize(), "Tone"]).size().unstack(fill_value=0)
             .reindex(columns=TONE, fill_value=0).reindex(hari, fill_value=0))
    x = list(range(len(hari)))
    bawah = [0] * len(hari)
    for t in TONE:
        ax.bar(x, tabel[t], bottom=bawah, color=D_TONE[t], width=0.72)
        bawah = [b + v for b, v in zip(bawah, tabel[t])]
    # label: tanggal 1, lalu kelipatan beberapa hari, dan tanggal terakhir
    maks_label = max(4, int(w_cm / 1.3))
    langkah = max(1, -(-len(hari) // maks_label))
    posisi = list(range(0, len(hari), langkah))
    terakhir = len(hari) - 1
    if posisi[-1] != terakhir:
        if terakhir - posisi[-1] >= langkah * 0.8:
            posisi.append(terakhir)       # cukup jauh: tambahkan label tanggal terakhir
        else:
            posisi[-1] = terakhir         # terlalu dekat: ganti label sebelumnya
    ax.set_xticks(posisi, [f"{hari[k].day} {BLN3[hari[k].month - 1]}" for k in posisi])
    ax.set_xlim(-0.6, len(hari) - 0.4)
    ax.yaxis.get_major_locator().set_params(integer=True)
    ax.set_ylim(0, max(1, max(bawah)) * 1.1)
    _gaya_ax(ax)
    _legenda(ax)
    return _png_tetap(fig, kiri=_kiri_cm(0.9, w_cm), atas=0.96, bawah=1.45 / h_cm)


def png_dash_donut(df, w_cm, h_cm):
    h = df["Tone"].value_counts().reindex(TONE, fill_value=0)
    total = int(h.sum())
    fig = plt.figure(figsize=(w_cm / 2.54, h_cm / 2.54))
    if total == 0:
        fig.text(0.5, 0.5, "Tidak ada data", ha="center", va="center", color=D_MUTED)
        return _png_tetap(fig)
    # donut di bagian atas, legenda di bawah; ukuran donut menyesuaikan ruang yang tersedia
    tinggi_legenda = 1.5 / h_cm
    sisi = min(w_cm * 0.82, h_cm * (1 - tinggi_legenda) * 0.92)
    ax = fig.add_axes([(1 - sisi / w_cm) / 2, tinggi_legenda + (1 - tinggi_legenda - sisi / h_cm) / 2,
                       sisi / w_cm, sisi / h_cm])
    ax.pie(h.values, colors=[D_TONE[t] for t in TONE], startangle=90, counterclock=False,
           wedgeprops=dict(width=0.38, edgecolor="white", linewidth=2))
    ax.text(0, 0.07, f"{total}", ha="center", va="center", fontsize=17, fontweight="bold", color=D_FG)
    ax.text(0, -0.2, "berita", ha="center", va="center", fontsize=7.5, color=D_MUTED)
    ax.set_aspect("equal")
    from matplotlib.lines import Line2D
    hd = [Line2D([0], [0], marker="o", color="w", markerfacecolor=D_TONE[t], markersize=6,
                 label=f"{t} {int(h[t])} ({h[t] / total:.0%})") for t in TONE]
    fig.legend(handles=hd, loc="lower center", bbox_to_anchor=(0.5, 0.0), ncol=1, frameon=False,
               fontsize=7.2, labelcolor=D_MUTED, handletextpad=0.3, labelspacing=0.35)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def png_dash_situs(df, w_cm, h_cm, n=8):
    s = df["Situs"].replace("", "(tanpa nama)").value_counts().head(n)[::-1]
    fig, ax = _cm(w_cm, h_cm)
    if s.empty:
        ax.axis("off")
        return _png_tetap(fig)
    ax.barh(range(len(s)), s.values, color=D_UNGU, height=0.62)
    ax.set_yticks(range(len(s)), [x if len(x) <= 22 else x[:21] + "…" for x in s.index])
    ax.xaxis.get_major_locator().set_params(integer=True)
    ax.set_xlim(0, max(s.values) * 1.12)
    for i, v in enumerate(s.values):
        ax.text(v + max(s.values) * 0.02, i, str(v), va="center", fontsize=7, color=D_MUTED)
    _gaya_ax(ax, "x")
    # ruang bawah sama dengan grafik "Tone per kata kunci" agar sumbu sejajar
    return _png_tetap(fig, kiri=_kiri_cm(3.0, w_cm), atas=0.97, bawah=1.45 / h_cm)


def png_dash_tone_topik(df, w_cm, h_cm):
    t = pd.crosstab(df["Topik"], df["Tone"]).reindex(columns=TONE, fill_value=0)
    t = t.loc[t.sum(axis=1).sort_values().index]
    fig, ax = _cm(w_cm, h_cm)
    if t.empty:
        ax.axis("off")
        return _png_tetap(fig)
    kiri = [0] * len(t)
    for tone in TONE:
        ax.barh(range(len(t)), t[tone], left=kiri, color=D_TONE[tone], height=0.55)
        for i, (k, v) in enumerate(zip(kiri, t[tone])):
            if v and v >= max(t.sum(axis=1)) * 0.08:
                ax.text(k + v / 2, i, str(v), ha="center", va="center", fontsize=7,
                        color="white", fontweight="bold")
        kiri = [a + b for a, b in zip(kiri, t[tone])]
    ax.set_yticks(range(len(t)), list(t.index))
    ax.xaxis.get_major_locator().set_params(integer=True)
    _gaya_ax(ax, "x")
    _legenda(ax)
    return _png_tetap(fig, kiri=_kiri_cm(2.4, w_cm), atas=0.97, bawah=1.45 / h_cm)


def kpi_dash(A):
    """Isi 5 kartu KPI: (label, nilai, sub-teks, teks perbandingan, warna)."""
    n = A["n"]
    pct = lambda x: f"{x / n:.0%} dari total" if n else "–"
    def banding(kini, lalu):
        if lalu is None:
            return ""
        d = kini - lalu
        return "sama dengan periode lalu" if d == 0 else f"{d:+d} vs {A['nama_lalu']}"
    situs = A["df"]["Situs"].replace("", pd.NA).dropna().nunique()
    isi = [("Total berita", n, "periode laporan", banding(n, A["n_lalu"]), D_UNGU)]
    for t in TONE:
        isi.append((t, A["tone"][t], pct(A["tone"][t]),
                    banding(A["tone"][t], A["tone_lalu"][t] if A["ada_lalu"] else None), D_TONE[t]))
    isi.append(("Situs media", situs, "media yang memberitakan", "", D_EMAS))
    return isi


# ===========================================================================
# PDF (reportlab)
# ===========================================================================

def buat_laporan_pdf(df, periode, df_lalu=None, nama_lalu=None,
                     judul="Laporan Pemantauan Berita SBSN", catatan="", lampiran=True,
                     rentang=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak,
                                    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    A = analisis(df, periode, df_lalu, nama_lalu)
    df = A["df"]
    LEBAR, TINGGI = A4
    MARGIN = 1.6 * cm
    ISI_W = LEBAR - 2 * MARGIN

    ss = getSampleStyleSheet()
    S_BAB = ParagraphStyle("bab", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=13,
                           textColor=colors.HexColor(UNGU), spaceBefore=10, spaceAfter=6)
    S_SUB = ParagraphStyle("sub", parent=S_BAB, fontSize=11, spaceBefore=8, spaceAfter=3,
                           textColor=colors.HexColor(PLUM))
    S_TEKS = ParagraphStyle("teks", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5,
                            leading=13.5, textColor=colors.HexColor("#2A2A2A"))
    S_LABEL = ParagraphStyle("lbl", parent=S_TEKS, fontName="Helvetica-Bold", fontSize=9,
                             textColor=colors.HexColor(UNGU), spaceBefore=4, spaceAfter=1)
    S_KECIL = ParagraphStyle("kecil", parent=S_TEKS, fontSize=7.8, leading=10)
    S_KECIL_P = ParagraphStyle("kecilp", parent=S_KECIL, textColor=colors.HexColor("#6B6B6B"))
    S_HEAD = ParagraphStyle("head", parent=S_KECIL, fontName="Helvetica-Bold",
                            textColor=colors.white, alignment=TA_CENTER)

    def img(png, lebar_cm):
        from reportlab.lib.utils import ImageReader
        ir = ImageReader(io.BytesIO(png))
        w, h = ir.getSize()
        return Image(io.BytesIO(png), width=lebar_cm * cm, height=lebar_cm * cm * h / w)

    def p(t, s=S_KECIL):
        return Paragraph(escape(str(t)), s)

    def h(t):
        return Paragraph(escape(t), S_HEAD)

    def tautan(judul_b, link):
        return Paragraph(f'<link href="{escape(link)}" color="#2A2A2A">{escape(judul_b)}</link>', S_KECIL)

    def bullets(items, gaya=S_TEKS):
        return ListFlowable([ListItem(Paragraph(escape(i), gaya), leftIndent=12, value="•")
                             for i in items], bulletType="bullet", start="•", leftIndent=12,
                            bulletColor=colors.HexColor(MAGENTA))

    def tabel(data, lebar, kolom_tone=None):
        t = Table(data, colWidths=lebar, repeatRows=1)
        g = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(UNGU)),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF6DA")]),
             ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor(EMAS)),
             ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D2DE")),
             ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
             ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        if kolom_tone is not None:
            for r, baris in enumerate(data[1:], start=1):
                tone = getattr(baris[kolom_tone], "text", "")
                if tone in TINT_TONE:
                    g.append(("BACKGROUND", (kolom_tone, r), (kolom_tone, r), colors.HexColor(TINT_TONE[tone])))
        t.setStyle(TableStyle(g))
        return t

    def kpi():
        isi = [("Total berita", A["n"], UNGU, A["n_lalu"])]
        isi += [(t, A["tone"][t], WARNA_TONE[t], A["tone_lalu"][t] if A["ada_lalu"] else None) for t in TONE]
        sel = []
        for label, nilai, warna, lalu in isi:
            dl = ""
            if lalu is not None:
                d = nilai - lalu
                dl = "sama dengan periode lalu" if d == 0 else f"{d:+d} vs {nama_lalu}"
            sel.append(Paragraph(
                f'<font color="#6B6B6B" size="7.5">{escape(label.upper())}</font><br/>'
                f'<font name="Helvetica-Bold" size="20" color="#2A2A2A">{nilai}</font>'
                + (f'<br/><font size="7.5" color="#6B6B6B">{escape(dl)}</font>' if dl else ""),
                ParagraphStyle("k", parent=S_TEKS, leading=17)))
        t = Table([sel], colWidths=[ISI_W / 4] * 4, rowHeights=[1.75 * cm])
        g = [("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 10)]
        for i, (_, _, warna, _) in enumerate(isi):
            g += [("LINEBEFORE", (i, 0), (i, 0), 4, colors.HexColor(warna)),
                  ("BACKGROUND", (i, 0), (i, 0), colors.HexColor("#F7F4F8"))]
        t.setStyle(TableStyle(g))
        return t

    S_PANEL = ParagraphStyle("panel", parent=S_TEKS, fontName="Helvetica-Bold", fontSize=9.5,
                             textColor=colors.HexColor(D_FG), spaceAfter=2)

    def judul_panel(t):
        return Paragraph(f'<font name="ZapfDingbats" size="7" color="{D_EMAS}">n</font>&nbsp;&nbsp;'
                         f'{escape(t)}', S_PANEL)

    def baris_dua(kiri, kanan, jarak=0.3 * cm):
        """Dua panel berdampingan dalam SATU baris tabel, sehingga tingginya selalu sama.
        kiri/kanan: (judul, png, lebar)."""
        (j1, g1, w1), (j2, g2, w2) = kiri, kanan
        t = Table([[judul_panel(j1), "", judul_panel(j2)],
                   [img(g1, w1 / cm - 0.4), "", img(g2, w2 / cm - 0.4)]],
                  colWidths=[w1, jarak, w2])
        gaya = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 1), (-1, 1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
                ("TOPPADDING", (0, 1), (-1, 1), 0), ("BOTTOMPADDING", (0, 1), (-1, 1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (1, 0), (1, -1), 0), ("RIGHTPADDING", (1, 0), (1, -1), 0)]
        for c in (0, 2):
            gaya += [("BOX", (c, 0), (c, 1), 0.6, colors.HexColor(D_LINE)),
                     ("BACKGROUND", (c, 0), (c, 1), colors.white)]
        t.setStyle(TableStyle(gaya))
        return t

    S_KLBL = ParagraphStyle("klbl", parent=S_TEKS, fontName="Helvetica-Bold", fontSize=6.8, leading=8,
                            textColor=colors.HexColor(D_MUTED), spaceAfter=6)
    S_KVAL = ParagraphStyle("kval", parent=S_TEKS, fontName="Helvetica-Bold", fontSize=20, leading=22,
                            textColor=colors.HexColor(D_FG), spaceAfter=4)
    S_KSUB = ParagraphStyle("ksub", parent=S_TEKS, fontSize=7, leading=9,
                            textColor=colors.HexColor(D_MUTED))

    def kartu_kpi():
        isi_k = kpi_dash(A)
        jarak = 0.25 * cm
        lebar = (ISI_W - jarak * (len(isi_k) - 1)) / len(isi_k)
        sel, kol, gaya = [], [], [("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]
        for k, (lbl, nilai, sub, banding, warna) in enumerate(isi_k):
            if k:
                sel.append("")
                kol.append(jarak)
            c = len(sel)
            isi_sel = [Paragraph(escape(lbl.upper()), S_KLBL), Paragraph(str(nilai), S_KVAL),
                       Paragraph(escape(sub), S_KSUB)]
            if banding:
                isi_sel.append(Paragraph(escape(banding), S_KSUB))
            sel.append(isi_sel)
            kol.append(lebar)
            gaya += [("BOX", (c, 0), (c, 0), 0.6, colors.HexColor(D_LINE)),
                     ("LINEBEFORE", (c, 0), (c, 0), 4, colors.HexColor(warna)),
                     ("BACKGROUND", (c, 0), (c, 0), colors.white),
                     ("LEFTPADDING", (c, 0), (c, 0), 9)]
        t = Table([sel], colWidths=kol)
        t.setStyle(TableStyle(gaya))
        return t

    isi = [Spacer(1, 2.6 * cm)]

    # 1. Ringkasan (susunan seperti dashboard: KPI -> tren + komposisi -> situs + tone per kata kunci)
    isi += [Paragraph("1. Ringkasan", S_BAB), kartu_kpi(), Spacer(1, 0.3 * cm)]
    w8 = (ISI_W - 0.3 * cm) * 8 / 12
    w4 = ISI_W - 0.3 * cm - w8
    w6 = (ISI_W - 0.3 * cm) / 2
    h1 = 6.2                                                  # tinggi acuan: tren harian
    isi.append(baris_dua(
        ("Tren harian per tone", png_dash_tren(df, w8 / cm - 0.4, h1, rentang), w8),
        ("Komposisi tone", png_dash_donut(df, w4 / cm - 0.4, h1), w4)))
    isi.append(Spacer(1, 0.3 * cm))
    h2 = max(4.2, 2.0 + 0.9 * df["Topik"].nunique(), 1.7 + 0.42 * min(8, df["Situs"].nunique()))
    isi.append(baris_dua(                                      # tinggi acuan: tone per kata kunci
        ("Situs teratas", png_dash_situs(df, w6 / cm - 0.4, h2), w6),
        ("Tone per kata kunci", png_dash_tone_topik(df, w6 / cm - 0.4, h2), w6)))
    isi += [Spacer(1, 0.3 * cm), p(A["ringkasan"], S_TEKS)]
    if catatan:
        isi.append(p(f"Catatan: {catatan}", S_KECIL_P))

    # 2. Insight utama
    isi.append(KeepTogether([Paragraph("2. Insight utama", S_BAB), bullets(A["insight_utama"])]))

    # 3. Pembahasan dan insight per topik
    # Setiap topik (3.1, 3.2, ...) dimulai di halaman baru
    isi += [PageBreak(), Paragraph("3. Pembahasan dan insight per topik", S_BAB)]
    for i, t in enumerate(A["topik"], start=1):
        if i > 1:
            isi.append(PageBreak())
        blok = [Paragraph(f"3.{i} {escape(t['nama'])}", S_SUB),
                p(f"{t['n']} berita · {_pct(t['share'])} dari total · indeks sentimen {t['indeks']:+d} "
                  f"({t['label']})" + (f" · periode lalu {t['n_lalu']} berita" if t["n_lalu"] is not None else ""),
                  S_KECIL_P),
                img(png_tone_topik(t), ISI_W / cm * 0.75),
                Paragraph("Pembahasan", S_LABEL), p(t["pembahasan"], S_TEKS)]
        isi.append(KeepTogether(blok))
        isi += [Paragraph("Insight", S_LABEL), bullets(t["insight"]),
                Paragraph("Rekomendasi", S_LABEL), bullets(t["rekomendasi"])]
        contoh = [("Positif", x) for x in t["contoh_pos"]] + [("Negatif", x) for x in t["contoh_neg"]]
        if contoh:
            data = [[h("Tone"), h("Contoh judul berita"), h("Media")]]
            for tone, (jd, st, lk) in contoh:
                data.append([p(tone), tautan(jd, lk), p(st)])
            isi += [Paragraph("Contoh berita", S_LABEL),
                    tabel(data, [1.8 * cm, ISI_W - 5.3 * cm, 3.5 * cm], kolom_tone=0)]
        isi.append(Spacer(1, 6))

    # Lampiran
    if lampiran and len(df):
        isi += [PageBreak(), Paragraph("Lampiran. Daftar berita", S_BAB),
                p("Judul berita dapat diklik pada versi digital untuk membuka sumbernya.", S_KECIL_P),
                Spacer(1, 4)]
        data = [[h("No"), h("Tanggal"), h("Media"), h("Judul"), h("Topik"), h("Tone"), h("Ritel")]]
        for i, r in enumerate(df.itertuples(), start=1):
            data.append([p(i), p(tgl_id(r.Tanggal, True)), p(r.Situs), tautan(r.Judul, r.Link),
                         p(r.Topik), p(r.Tone), p(r.Ritel)])
        isi.append(tabel(data, [0.9 * cm, 1.9 * cm, 2.6 * cm, ISI_W - 11.1 * cm, 2.3 * cm,
                                1.5 * cm, 1.9 * cm], kolom_tone=5))

    dicetak = datetime.now()

    def kaki(c, doc):
        c.setStrokeColor(colors.HexColor(MAGENTA))
        c.setLineWidth(0.8)
        c.line(MARGIN, 1.3 * cm, LEBAR - MARGIN, 1.3 * cm)
        c.setFillColor(colors.HexColor("#6B6B6B"))
        c.setFont("Helvetica", 7.5)
        c.drawString(MARGIN, 0.9 * cm, f"{UNIT} · dicetak {tgl_id(dicetak)} {dicetak:%H:%M}")
        c.drawRightString(LEBAR - MARGIN, 0.9 * cm, f"Halaman {doc.page}")

    def pertama(c, doc):
        c.saveState()
        tg = 3.3 * cm
        c.setFillColor(colors.HexColor(UNGU))
        c.rect(0, TINGGI - tg, LEBAR, tg, stroke=0, fill=1)
        c.setFillColor(colors.HexColor(MAGENTA))
        pth = c.beginPath()
        pth.moveTo(LEBAR * 0.72, TINGGI)
        pth.lineTo(LEBAR, TINGGI)
        pth.lineTo(LEBAR, TINGGI - tg)
        pth.lineTo(LEBAR * 0.86, TINGGI - tg)
        pth.close()
        c.drawPath(pth, stroke=0, fill=1)
        c.setFillColor(colors.HexColor(EMAS))
        c.rect(0, TINGGI - tg - 0.18 * cm, LEBAR, 0.18 * cm, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 18)
        c.drawString(MARGIN, TINGGI - 1.45 * cm, judul.upper())
        c.setFont("Helvetica", 11)
        c.drawString(MARGIN, TINGGI - 2.15 * cm, f"Periode: {periode}")
        c.setFont("Helvetica", 8.5)
        c.drawString(MARGIN, TINGGI - 2.7 * cm, UNIT)
        c.setFont("Helvetica-Bold", 12)
        c.drawRightString(LEBAR - MARGIN, TINGGI - 1.45 * cm, "DJPPR")
        kaki(c, doc)
        c.restoreState()

    def lanjut(c, doc):
        c.saveState()
        c.setFillColor(colors.HexColor(UNGU))
        c.rect(0, TINGGI - 0.9 * cm, LEBAR, 0.9 * cm, stroke=0, fill=1)
        c.setFillColor(colors.HexColor(EMAS))
        c.rect(0, TINGGI - 0.98 * cm, LEBAR, 0.08 * cm, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 9)
        c.drawString(MARGIN, TINGGI - 0.6 * cm, judul)
        c.setFont("Helvetica", 9)
        c.drawRightString(LEBAR - MARGIN, TINGGI - 0.6 * cm, periode)
        kaki(c, doc)
        c.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=MARGIN, rightMargin=MARGIN,
                            topMargin=1.5 * cm, bottomMargin=1.7 * cm,
                            title=f"{judul} - {periode}", author=UNIT)
    doc.build(isi, onFirstPage=pertama, onLaterPages=lanjut)
    return buf.getvalue()


# ===========================================================================
# WORD (python-docx)
# ===========================================================================

def buat_laporan_docx(df, periode, df_lalu=None, nama_lalu=None,
                      judul="Laporan Pemantauan Berita SBSN", catatan="", lampiran=True,
                     rentang=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    A = analisis(df, periode, df_lalu, nama_lalu)
    df = A["df"]

    def rgb(hx):
        return RGBColor.from_string(hx.lstrip("#").upper())

    def arsir(sel, hx):
        tcPr = sel._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hx.lstrip("#").upper())
        tcPr.append(shd)

    def batas_kiri(sel, hx, tebal=24):
        tcPr = sel._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        kiri = OxmlElement("w:left")
        kiri.set(qn("w:val"), "single")
        kiri.set(qn("w:sz"), str(tebal))
        kiri.set(qn("w:color"), hx.lstrip("#").upper())
        b.append(kiri)
        tcPr.append(b)

    def hyperlink(par, teks, url, ukuran=8):
        r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                  is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(ukuran * 2))
        rPr.append(sz)
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "2A2A2A")
        rPr.append(col)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = teks
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        hl.append(run)
        par._p.append(hl)

    def teks_sel(sel, teks, ukuran=8, tebal=False, warna=None, rata=None):
        sel.text = ""
        par = sel.paragraphs[0]
        r = par.add_run(str(teks))
        r.font.size = Pt(ukuran)
        r.bold = tebal
        if warna:
            r.font.color.rgb = rgb(warna)
        if rata:
            par.alignment = rata
        par.paragraph_format.space_after = Pt(0)
        return par

    def tabel(header, baris, lebar_cm, kolom_tone=None, kolom_link=None):
        t = doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for i, hd in enumerate(header):
            sel = t.rows[0].cells[i]
            arsir(sel, UNGU)
            teks_sel(sel, hd, 8, True, "#FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        # header berulang di tiap halaman
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for n, b in enumerate(baris):
            sel = t.add_row().cells
            for i, v in enumerate(b[:len(header)]):
                if kolom_link is not None and i == kolom_link[0] and b[kolom_link[1]]:
                    sel[i].text = ""
                    par = sel[i].paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    hyperlink(par, str(v), b[kolom_link[1]])
                else:
                    teks_sel(sel[i], v)
                if n % 2:
                    arsir(sel[i], "#FFF6DA")
            if kolom_tone is not None and b[kolom_tone] in TINT_TONE:
                arsir(sel[kolom_tone], TINT_TONE[b[kolom_tone]])
        atur_lebar(t, lebar_cm)
        return t

    def atur_lebar(t, lebar_cm):
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(lebar_cm):
                gc.set(qn("w:w"), str(int(lebar_cm[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(lebar_cm):
                row.cells[i].width = Cm(w)

    def bab(teks):
        hd = doc.add_heading(teks, level=1)
        for r in hd.runs:
            r.font.color.rgb = rgb(UNGU)
            r.font.size = Pt(14)
        return hd

    def sub(teks):
        hd = doc.add_heading(teks, level=2)
        for r in hd.runs:
            r.font.color.rgb = rgb(PLUM)
            r.font.size = Pt(12)
        return hd

    def label(teks):
        par = doc.add_paragraph()
        r = par.add_run(teks)
        r.bold = True
        r.font.color.rgb = rgb(UNGU)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.space_before = Pt(4)

    def poin(items):
        for i in items:
            doc.add_paragraph(i, style="List Bullet")

    def gambar(png, lebar_cm):
        doc.add_picture(io.BytesIO(png), width=Cm(lebar_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc = Document()
    sect = doc.sections[0]
    sect.page_width, sect.page_height = Cm(21), Cm(29.7)
    sect.left_margin = sect.right_margin = Cm(1.8)
    sect.top_margin, sect.bottom_margin = Cm(1.6), Cm(1.6)
    ISI = 21 - 3.6

    gaya = doc.styles["Normal"]
    gaya.font.name = "Arial"
    gaya.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    gaya.font.size = Pt(10)
    for nama_gaya in ["Heading 1", "Heading 2"]:
        doc.styles[nama_gaya].font.name = "Arial"

    # Footer: unit + nomor halaman
    fp = sect.footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(f"{UNIT} · dicetak {tgl_id(datetime.now())} · Halaman ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = rgb("#6B6B6B")
    for kode in ["begin", None, "end"]:
        run = fp.add_run()
        run.font.size = Pt(7.5)
        if kode:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kode)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = "PAGE"
            run._r.append(it)

    # Kepala laporan (tabel berwarna)
    kepala = doc.add_table(rows=1, cols=1)
    sel = kepala.rows[0].cells[0]
    arsir(sel, UNGU)
    sel.text = ""
    p1 = sel.paragraphs[0]
    r = p1.add_run(judul.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), rgb("#FFFFFF")
    p2 = sel.add_paragraph()
    r = p2.add_run(f"Periode: {periode}")
    r.font.size, r.font.color.rgb = Pt(11), rgb("#FFFFFF")
    p3 = sel.add_paragraph()
    r = p3.add_run(UNIT)
    r.font.size, r.font.color.rgb = Pt(8.5), rgb("#FFE9A8")
    garis = doc.add_table(rows=1, cols=1)
    arsir(garis.rows[0].cells[0], EMAS)
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    garis.rows[0].height = Cm(0.15)
    garis.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    teks_sel(garis.rows[0].cells[0], "", 2)

    def batas_sel(sel, hx="E3DCE8", tebal=6, kiri=None):
        tcPr = sel._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for sisi in ["top", "left", "bottom", "right"]:
            e = OxmlElement(f"w:{sisi}")
            e.set(qn("w:val"), "single")
            if sisi == "left" and kiri:
                e.set(qn("w:sz"), "32")
                e.set(qn("w:color"), kiri.lstrip("#").upper())
            else:
                e.set(qn("w:sz"), str(tebal))
                e.set(qn("w:color"), hx)
            b.append(e)
        tcPr.append(b)

    def tanpa_batas(t):
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for sisi in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            e = OxmlElement(f"w:{sisi}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        tblPr.append(b)

    def isi_panel(sel, judul_p, png, lebar_cm):
        batas_sel(sel)
        sel.text = ""
        par = sel.paragraphs[0]
        r = par.add_run("■ ")
        r.font.color.rgb = rgb(D_EMAS)
        r.font.size = Pt(8)
        r = par.add_run(judul_p)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), rgb(D_FG)
        par2 = sel.add_paragraph()
        par2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par2.add_run().add_picture(io.BytesIO(png), width=Cm(lebar_cm))

    # 1. Ringkasan (susunan seperti dashboard)
    bab("1. Ringkasan")
    isi_k = kpi_dash(A)
    k = doc.add_table(rows=1, cols=len(isi_k))
    tanpa_batas(k)
    atur_lebar(k, [ISI / len(isi_k)] * len(isi_k))
    for i, (lbl, nilai, sub_t, banding, warna) in enumerate(isi_k):
        sel = k.rows[0].cells[i]
        batas_sel(sel, kiri=warna)
        teks_sel(sel, lbl.upper(), 7, True, D_MUTED)
        pn = sel.add_paragraph()
        pn.paragraph_format.space_before = Pt(5)        # jarak label -> angka
        pn.paragraph_format.space_after = Pt(3)
        r = pn.add_run(str(nilai))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(20), rgb(D_FG)
        for teks in [sub_t, banding]:
            if teks:
                ps = sel.add_paragraph()
                ps.paragraph_format.space_after = Pt(0)
                rr = ps.add_run(teks)
                rr.font.size, rr.font.color.rgb = Pt(7), rgb(D_MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    w8, w4 = ISI * 8 / 12, ISI * 4 / 12
    h1 = 6.2                                                   # tinggi acuan: tren harian
    g1 = doc.add_table(rows=1, cols=2)
    tanpa_batas(g1)
    atur_lebar(g1, [w8, w4])
    isi_panel(g1.rows[0].cells[0], "Tren harian per tone", png_dash_tren(df, w8 - 0.6, h1, rentang),
              w8 - 0.6)
    isi_panel(g1.rows[0].cells[1], "Komposisi tone", png_dash_donut(df, w4 - 0.6, h1), w4 - 0.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    w6 = ISI / 2
    h2 = max(4.2, 2.0 + 0.9 * df["Topik"].nunique(), 1.7 + 0.42 * min(8, df["Situs"].nunique()))
    g2 = doc.add_table(rows=1, cols=2)
    tanpa_batas(g2)
    atur_lebar(g2, [w6, w6])
    isi_panel(g2.rows[0].cells[0], "Situs teratas", png_dash_situs(df, w6 - 0.6, h2), w6 - 0.6)
    isi_panel(g2.rows[0].cells[1], "Tone per kata kunci", png_dash_tone_topik(df, w6 - 0.6, h2),
              w6 - 0.6)
    doc.add_paragraph()
    doc.add_paragraph(A["ringkasan"])
    if catatan:
        cp = doc.add_paragraph().add_run(f"Catatan: {catatan}")
        cp.font.size, cp.italic = Pt(8.5), True

    # 2. Insight utama
    bab("2. Insight utama")
    poin(A["insight_utama"])

    # 3. Pembahasan dan insight per topik
    # Setiap topik (3.1, 3.2, ...) dimulai di halaman baru
    doc.add_page_break()
    bab("3. Pembahasan dan insight per topik")
    for i, t in enumerate(A["topik"], start=1):
        if i > 1:
            doc.add_page_break()
        sub(f"3.{i} {t['nama']}")
        info = doc.add_paragraph().add_run(
            f"{t['n']} berita · {_pct(t['share'])} dari total · indeks sentimen {t['indeks']:+d} "
            f"({t['label']})" + (f" · periode lalu {t['n_lalu']} berita" if t["n_lalu"] is not None else ""))
        info.font.size, info.font.color.rgb = Pt(8.5), rgb("#6B6B6B")
        gambar(png_tone_topik(t), ISI * 0.75)
        label("Pembahasan")
        doc.add_paragraph(t["pembahasan"])
        label("Insight")
        poin(t["insight"])
        label("Rekomendasi")
        poin(t["rekomendasi"])
        contoh = [("Positif", x) for x in t["contoh_pos"]] + [("Negatif", x) for x in t["contoh_neg"]]
        if contoh:
            label("Contoh berita")
            tabel(["Tone", "Contoh judul berita", "Media"],
                  [[tone, jd, st, lk] for tone, (jd, st, lk) in contoh],
                  [2.0, ISI - 5.5, 3.5], kolom_tone=0, kolom_link=(1, 3))

    # Lampiran
    if lampiran and len(df):
        doc.add_page_break()
        bab("Lampiran. Daftar berita")
        tabel(["No", "Tanggal", "Media", "Judul", "Topik", "Tone", "Ritel"],
              [[i, tgl_id(r.Tanggal, True), r.Situs, r.Judul, r.Topik, r.Tone, r.Ritel, r.Link]
               for i, r in enumerate(df.itertuples(), start=1)],
              [0.9, 2.0, 2.6, ISI - 11.3, 2.3, 1.6, 1.9], kolom_tone=5, kolom_link=(3, 7))

    doc.core_properties.title = f"{judul} - {periode}"
    doc.core_properties.author = UNIT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Pakai: python laporan.py <file_excel_hasil>")
        sys.exit(1)
    data = pd.read_excel(sys.argv[1])
    t = pd.to_datetime(data["Tanggal"], errors="coerce").dropna()
    periode = f"{tgl_id(t.min())} – {tgl_id(t.max())}" if len(t) else "-"
    dasar = sys.argv[1].rsplit(".", 1)[0] + "_laporan"
    with open(dasar + ".pdf", "wb") as f:
        f.write(buat_laporan_pdf(data, periode))
    with open(dasar + ".docx", "wb") as f:
        f.write(buat_laporan_docx(data, periode))
    print(f"Laporan tersimpan: {dasar}.pdf dan {dasar}.docx")
